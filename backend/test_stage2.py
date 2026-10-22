from fastapi.testclient import TestClient
import openpyxl
import io
from main import app

client = TestClient(app)

print("--- TEST 1: ENTITY RESOLUTION ---")
res_ent = client.get("/api/etl/entities")
print("Status:", res_ent.status_code, "Total banks:", res_ent.json()["total_banks"])
assert res_ent.status_code == 200
assert res_ent.json()["total_banks"] == 21

print("\n--- TEST 2: POST /api/etl/upload ---")
res_etl = client.post("/api/etl/upload", data={"use_sample_datasets": "true", "stage": "EARLY"})
print("ETL Status:", res_etl.status_code)
print("ETL Summary:", res_etl.json()["evaluation_summary"])
assert res_etl.status_code == 200

print("\n--- TEST 3: VERIFY MATRIX METRICS ---")
res_mat = client.get("/api/matrix")
matrix_data = {item["masked_alias"]: item for item in res_mat.json()}

bank_a = matrix_data["Bank A"]
bank_b = matrix_data["Bank B"]
bank_d = matrix_data["Bank D"]
bank_n = matrix_data["Bank N"]
bank_r = matrix_data["Bank R"]

print(f"Bank A (B2 Kuotasi SRBI):  raw={bank_a['raw_metrics']['B2']}%, passed={bank_a['passed_flags']['B2']}, score={bank_a['weighted_scores']['B2']}")
print(f"Bank B (B3 Kuotasi OIS):   raw={bank_b['raw_metrics']['B3']}%, passed={bank_b['passed_flags']['B3']}, score={bank_b['weighted_scores']['B3']}")
print(f"Bank D (B6 Jual SRBI Sek): raw={bank_d['raw_metrics']['B6']}%, passed={bank_d['passed_flags']['B6']}, score={bank_d['weighted_scores']['B6']}")
print(f"Bank N (B4 Line Repo):     raw={bank_n['raw_metrics']['B4']}%, passed={bank_n['passed_flags']['B4']}, score={bank_n['weighted_scores']['B4']}")
print(f"Bank R (B7a Repo Sekunder):raw={bank_r['raw_metrics']['B7a']}%, passed={bank_r['passed_flags']['B7a']}, score={bank_r['weighted_scores']['B7a']}")

assert bank_a["raw_metrics"]["B2"] == 76.67
assert bank_b["raw_metrics"]["B3"] == 21.67
assert bank_d["raw_metrics"]["B6"] == 65.0
assert bank_n["raw_metrics"]["B4"] == 45.0
assert bank_r["raw_metrics"]["B7a"] == 72.22

print("\n--- TEST 4: DYNAMIC THRESHOLD RECALCULATE ---")
res_recalc_inter = client.post("/api/evaluation/recalculate?stage=INTERMEDIATE")
print("Intermediate Result:", res_recalc_inter.json()["data"])
assert res_recalc_inter.status_code == 200

res_recalc_early = client.post("/api/evaluation/recalculate?stage=EARLY")
print("Early Result:", res_recalc_early.json()["data"])
assert res_recalc_early.status_code == 200

print("\n--- TEST 5: GET /api/export/worksheet/{bank_id} ---")
for b_code in ["DU-01", "DU-18"]:
    res_ws = client.get(f"/api/export/worksheet/{b_code}")
    print(f"Export {b_code} status:", res_ws.status_code, "Content length:", len(res_ws.content))
    assert res_ws.status_code == 200
    
    wb = openpyxl.load_workbook(io.BytesIO(res_ws.content))
    print(f"  Sheets for {b_code}:", wb.sheetnames)
    assert "Ringkasan_Evaluasi" in wb.sheetnames
    assert "Daftar_Temuan" in wb.sheetnames
    assert "Transaksi_Gugur" in wb.sheetnames
    
    # Check Sheet 3 content
    ws3 = wb["Transaksi_Gugur"]
    rows = list(ws3.iter_rows(values_only=True))
    print(f"  Rows in Transaksi_Gugur for {b_code}: {len(rows)}")
    if b_code == "DU-18":
        assert len(rows) >= 13  # 12 transactions + header
        print("  Sample row Bank R (DU-18):", rows[4])

print("\nALL STAGE 2 TESTS PASSED PERFECTLY!")
