"""
Auto-Splitter Worksheet Per Bank
Menghasilkan workbook Excel individual (.xlsx) per bank DU PUVA sesuai standar Bank Indonesia.
"""

import io
import json
from datetime import datetime
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from models import MasterBank, EvaluationPeriod, BankScoreMatrix, FindingTicket
from services.rule_engine import THRESHOLDS, WEIGHTS, PARAMETER_NAMES, PARAMETER_DETAILS


def generate_bank_worksheet(db: Session, bank_id: str) -> io.BytesIO:
    """
    Generate file Excel (.xlsx) rapi dan terformat untuk 1 bank DU PUVA:
    1. Sheet 'Ringkasan_Evaluasi': Profil bank, skor PUVA B1-B8 & OM, serta komposit.
    2. Sheet 'Daftar_Temuan': Tiket Finding_ID dan status konfirmasi.
    3. Sheet 'Transaksi_Gugur': Rincian baris transaksi mentah yang gugur.
    """
    bank = db.query(MasterBank).filter(MasterBank.Bank_ID == bank_id).first()
    if not bank:
        # Fallback cari berdasarkan alias misal 'Bank A'
        bank = db.query(MasterBank).filter(MasterBank.Masked_Alias == bank_id).first()
    if not bank:
        raise ValueError(f"Bank dengan ID '{bank_id}' tidak ditemukan.")

    period = db.query(EvaluationPeriod).order_by(EvaluationPeriod.id.desc()).first()
    stage = period.Active_Threshold_Stage if period else "EARLY"
    triwulan = period.Triwulan if period else "2026-Q2"

    matrix = db.query(BankScoreMatrix).filter(
        BankScoreMatrix.Bank_ID == bank.Bank_ID
    ).order_by(BankScoreMatrix.id.desc()).first()

    tickets = db.query(FindingTicket).filter(
        FindingTicket.Bank_ID == bank.Bank_ID
    ).all()

    # Inisialisasi Workbook openpyxl
    wb = Workbook()
    # Hapus sheet default
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    emerald_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    light_emerald_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    light_rose_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    light_amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    font_title = Font(name="Calibri", size=14, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_pass = Font(name="Calibri", size=10, bold=True, color="047857")
    font_fail = Font(name="Calibri", size=10, bold=True, color="BE123C")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="0F172A"))

    # ==========================================================================
    # SHEET 1: RINGKASAN EVALUASI
    # ==========================================================================
    ws1 = wb.create_sheet(title="Ringkasan_Evaluasi")
    ws1.views.sheetView[0].showGridLines = True

    # Header Dokumen
    ws1.cell(row=1, column=1, value="BANK INDONESIA - DEPARTEMEN PENGEMBANGAN PASAR KEUANGAN (DPPK)").font = font_subtitle
    ws1.cell(row=2, column=1, value=f"LEMBAR KONFIRMASI EVALUASI DEALER UTAMA PUVA - {triwulan}").font = font_title
    ws1.cell(row=3, column=1, value=f"Tahap Evaluasi: {stage} STATE | Dokumen Resmi Hasil Penilaian Kepatuhan PADG No. 23/2024").font = font_subtitle

    # Profil Bank
    ws1.cell(row=5, column=1, value="ID Bank:").font = font_bold
    ws1.cell(row=5, column=2, value=bank.Bank_ID).font = font_regular
    ws1.cell(row=5, column=4, value="Kelompok KBMI:").font = font_bold
    ws1.cell(row=5, column=5, value=bank.Kelompok_KBMI).font = font_regular

    ws1.cell(row=6, column=1, value="Masked Alias:").font = font_bold
    ws1.cell(row=6, column=2, value=bank.Masked_Alias).font = font_regular
    ws1.cell(row=6, column=4, value="Kode Bloomberg:").font = font_bold
    ws1.cell(row=6, column=5, value=bank.Kode_Bloomberg).font = font_regular

    ws1.cell(row=7, column=1, value="Nama Resmi:").font = font_bold
    ws1.cell(row=7, column=2, value=bank.Nama_Resmi_Bank).font = font_regular
    ws1.cell(row=7, column=4, value="Kode SPPA/Antasena:").font = font_bold
    ws1.cell(row=7, column=5, value=bank.Kode_SPPA_Antasena).font = font_regular

    # Tabel Penilaian Komponen PUVA (60%)
    ws1.cell(row=9, column=1, value="A. KOMPONEN PASAR UANG & PASAR VALAS (PUVA - MAKS 60%)").font = font_bold

    headers_puva = ["Kode", "Parameter Kewajiban", "Kriteria PADG No. 23/2024", "Capaian (%)", "Threshold", "Status", "Bobot", "Skor Bank"]
    for col_idx, h_text in enumerate(headers_puva, start=1):
        c = ws1.cell(row=10, column=col_idx, value=h_text)
        c.fill = navy_fill
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    puva_keys = [
        ("B1", "b1_raw", "b1_passed", "b1_score"),
        ("B2", "b2_raw", "b2_passed", "b2_score"),
        ("B3", "b3_raw", "b3_passed", "b3_score"),
        ("B4", "b4_raw", "b4_passed", "b4_score"),
        ("B5", "b5_raw", "b5_passed", "b5_score"),
        ("B6", "b6_raw", "b6_passed", "b6_score"),
        ("B7a", "b7a_raw", "b7a_passed", "b7a_score"),
        ("B7b", "b7b_raw", "b7b_passed", "b7b_score"),
        ("B8", "b8_raw", "b8_passed", "b8_score"),
    ]

    curr_row = 11
    stage_thresholds = THRESHOLDS.get(stage, THRESHOLDS["EARLY"])

    for code, raw_col, pass_col, score_col in puva_keys:
        raw_val = getattr(matrix, raw_col, 0.0) if matrix else 0.0
        is_pass = getattr(matrix, pass_col, False) if matrix else False
        score_val = getattr(matrix, score_col, 0.0) if matrix else 0.0
        th_val = stage_thresholds.get(code, 80.0)
        weight_val = WEIGHTS.get(code, 0.0)

        ws1.cell(row=curr_row, column=1, value=code).alignment = Alignment(horizontal="center")
        ws1.cell(row=curr_row, column=2, value=PARAMETER_NAMES.get(code, code))
        ws1.cell(row=curr_row, column=3, value=PARAMETER_DETAILS.get(code, "-"))
        
        c_raw = ws1.cell(row=curr_row, column=4, value=f"{raw_val:.2f}%")
        c_raw.alignment = Alignment(horizontal="right")

        c_th = ws1.cell(row=curr_row, column=5, value=f">= {th_val:.1f}%")
        c_th.alignment = Alignment(horizontal="center")

        c_status = ws1.cell(row=curr_row, column=6, value="MEMENUHI" if is_pass else "TEMUAN")
        c_status.alignment = Alignment(horizontal="center")
        if is_pass:
            c_status.fill = light_emerald_fill
            c_status.font = font_pass
        else:
            c_status.fill = light_rose_fill
            c_status.font = font_fail

        c_w = ws1.cell(row=curr_row, column=7, value=f"{weight_val:.1f}%")
        c_w.alignment = Alignment(horizontal="right")

        c_sc = ws1.cell(row=curr_row, column=8, value=f"{score_val:.2f}%")
        c_sc.alignment = Alignment(horizontal="right")
        c_sc.font = font_bold

        for col_idx in range(1, 9):
            ws1.cell(row=curr_row, column=col_idx).border = cell_border
        curr_row += 1

    # Subtotal PUVA
    puva_total = getattr(matrix, "total_puva_score", 0.0) if matrix else 0.0
    ws1.cell(row=curr_row, column=2, value="SUBTOTAL KOMPONEN PUVA").font = font_bold
    ws1.cell(row=curr_row, column=7, value="60.0%").font = font_bold
    c_tot_puva = ws1.cell(row=curr_row, column=8, value=f"{puva_total:.2f}%")
    c_tot_puva.font = font_bold
    c_tot_puva.alignment = Alignment(horizontal="right")
    for col_idx in range(1, 9):
        ws1.cell(row=curr_row, column=col_idx).border = cell_border
        ws1.cell(row=curr_row, column=col_idx).fill = gray_fill
    curr_row += 2

    # Tabel Penilaian Komponen OM (40%)
    ws1.cell(row=curr_row, column=1, value="B. KOMPONEN OPERASI MONETER (OM - MAKS 40%)").font = font_bold
    curr_row += 1

    for col_idx, h_text in enumerate(headers_puva, start=1):
        c = ws1.cell(row=curr_row, column=col_idx, value=h_text)
        c.fill = navy_fill
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center")
    curr_row += 1

    om_keys = [
        ("OM1", "om1_raw", "om1_passed", "om1_score"),
        ("OM2", "om2_raw", "om2_passed", "om2_score"),
        ("OM3", "om3_raw", "om3_passed", "om3_score"),
    ]

    for code, raw_col, pass_col, score_col in om_keys:
        raw_val = getattr(matrix, raw_col, 100.0) if matrix else 100.0
        is_pass = getattr(matrix, pass_col, True) if matrix else True
        score_val = getattr(matrix, score_col, 0.0) if matrix else 0.0
        th_val = stage_thresholds.get(code, 60.0)
        weight_val = WEIGHTS.get(code, 0.0)

        ws1.cell(row=curr_row, column=1, value=code).alignment = Alignment(horizontal="center")
        ws1.cell(row=curr_row, column=2, value=PARAMETER_NAMES.get(code, code))
        ws1.cell(row=curr_row, column=3, value=PARAMETER_DETAILS.get(code, "-"))
        ws1.cell(row=curr_row, column=4, value=f"{raw_val:.2f}%").alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=5, value=f">= {th_val:.1f}%").alignment = Alignment(horizontal="center")
        
        c_status = ws1.cell(row=curr_row, column=6, value="MEMENUHI" if is_pass else "TEMUAN")
        c_status.alignment = Alignment(horizontal="center")
        c_status.font = font_pass if is_pass else font_fail
        
        ws1.cell(row=curr_row, column=7, value=f"{weight_val:.1f}%").alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=8, value=f"{score_val:.2f}%").alignment = Alignment(horizontal="right")

        for col_idx in range(1, 9):
            ws1.cell(row=curr_row, column=col_idx).border = cell_border
        curr_row += 1

    # Subtotal OM
    om_total = getattr(matrix, "total_om_score", 40.0) if matrix else 40.0
    ws1.cell(row=curr_row, column=2, value="SUBTOTAL KOMPONEN OM").font = font_bold
    ws1.cell(row=curr_row, column=7, value="40.0%").font = font_bold
    c_tot_om = ws1.cell(row=curr_row, column=8, value=f"{om_total:.2f}%")
    c_tot_om.font = font_bold
    c_tot_om.alignment = Alignment(horizontal="right")
    for col_idx in range(1, 9):
        ws1.cell(row=curr_row, column=col_idx).border = cell_border
        ws1.cell(row=curr_row, column=col_idx).fill = gray_fill
    curr_row += 2

    # Ringkasan Total & Kepatuhan
    comp_score = getattr(matrix, "composite_score", 100.0) if matrix else 100.0
    compliance = getattr(matrix, "compliance_status", "MEMENUHI") if matrix else "MEMENUHI"

    ws1.cell(row=curr_row, column=1, value="C. KESIMPULAN EVALUASI & TINDAK LANJUT").font = font_bold
    curr_row += 1

    ws1.cell(row=curr_row, column=1, value="Total Skor Komposit (100%):").font = font_bold
    c_comp = ws1.cell(row=curr_row, column=2, value=f"{comp_score:.2f}%")
    c_comp.font = Font(name="Calibri", size=12, bold=True, color="0F172A")

    curr_row += 1
    ws1.cell(row=curr_row, column=1, value="Status Evaluasi Akhir:").font = font_bold
    c_status_final = ws1.cell(row=curr_row, column=2, value=compliance)
    c_status_final.font = font_bold
    if compliance == "MEMENUHI":
        c_status_final.fill = light_emerald_fill
        c_status_final.font = font_pass
    elif compliance == "PEMBINAAN_DSMM":
        c_status_final.fill = light_amber_fill
        c_status_final.font = Font(name="Calibri", size=11, bold=True, color="B45309")
    else:
        c_status_final.fill = light_rose_fill
        c_status_final.font = font_fail

    curr_row += 1
    ws1.cell(row=curr_row, column=1, value="Tindak Lanjut BI:").font = font_bold
    if comp_score >= 100.0:
        action_text = "Memenuhi seluruh kewajiban Dealer Utama PUVA (Lulus Penuh)"
    elif comp_score >= 90.0:
        action_text = "Surat/Email Pembinaan oleh Departemen Surveilans & Manajemen Risiko (DSMM)"
    else:
        action_text = "Sanksi Administratif berupa Surat Teguran Tertulis"
    ws1.cell(row=curr_row, column=2, value=action_text).font = font_regular

    # Atur lebar kolom ws1
    col_widths_ws1 = [8, 30, 48, 14, 14, 16, 12, 14]
    for i, w in enumerate(col_widths_ws1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ==========================================================================
    # SHEET 2: DAFTAR TEMUAN & SANGGAHAN
    # ==========================================================================
    ws2 = wb.create_sheet(title="Daftar_Temuan")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=1, column=1, value=f"DAFTAR TIKET TEMUAN & REKONSILIASI SANGGAHAN - {bank.Nama_Resmi_Bank}").font = font_title
    ws2.cell(row=2, column=1, value="Silakan lengkapi formulir sanggahan dan lampiran file bukti bila terdapat ketidaksesuaian data.").font = font_subtitle

    headers_findings = [
        "Finding ID", "Parameter", "Nilai Rekap BI", "Klaim Bank",
        "Alasan Sanggahan Bank", "File Bukti", "Rekomendasi AI", "Status Keputusan BI"
    ]
    for col_idx, h_text in enumerate(headers_findings, start=1):
        c = ws2.cell(row=4, column=col_idx, value=h_text)
        c.fill = navy_fill
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    curr_row_2 = 5
    if tickets:
        for t in tickets:
            ws2.cell(row=curr_row_2, column=1, value=t.Finding_ID).font = font_bold
            ws2.cell(row=curr_row_2, column=2, value=f"{t.Parameter_Code} - {PARAMETER_NAMES.get(t.Parameter_Code, '')}")
            ws2.cell(row=curr_row_2, column=3, value=t.Nilai_Rekap_BI or "-")
            ws2.cell(row=curr_row_2, column=4, value=t.Nilai_Klaim_Bank or "[Belum Diisi]")
            ws2.cell(row=curr_row_2, column=5, value=t.Alasan_Sanggahan or "[Belum Diisi]")
            ws2.cell(row=curr_row_2, column=6, value=t.File_Bukti_Name or "-")
            ws2.cell(row=curr_row_2, column=7, value=t.AI_Recommendation_Status or "-")
            
            c_dec = ws2.cell(row=curr_row_2, column=8, value=t.BI_Decision_Status or "OPEN")
            c_dec.alignment = Alignment(horizontal="center")
            c_dec.font = font_bold
            if t.BI_Decision_Status == "APPROVED":
                c_dec.fill = light_emerald_fill
                c_dec.font = font_pass
            elif t.BI_Decision_Status == "REJECTED":
                c_dec.fill = light_rose_fill
                c_dec.font = font_fail

            for col_idx in range(1, 9):
                ws2.cell(row=curr_row_2, column=col_idx).border = cell_border
            curr_row_2 += 1
    else:
        ws2.cell(row=curr_row_2, column=1, value="Tidak ada temuan. Bank memenuhi seluruh parameter evaluasi.").font = font_pass
        ws2.merge_cells(start_row=curr_row_2, start_column=1, end_row=curr_row_2, end_column=8)
        for col_idx in range(1, 9):
            ws2.cell(row=curr_row_2, column=col_idx).border = cell_border

    col_widths_ws2 = [24, 25, 25, 25, 45, 28, 22, 18]
    for i, w in enumerate(col_widths_ws2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ==========================================================================
    # SHEET 3: TRANSAKSI GUGUR (DETAIL DISKUALIFIKASI ATURAN PADG)
    # ==========================================================================
    ws3 = wb.create_sheet(title="Transaksi_Gugur")
    ws3.views.sheetView[0].showGridLines = True

    ws3.cell(row=1, column=1, value=f"RINCIAN TRANSAKSI / KUOTASI GUGUR (TIDAK MEMENUHI SYARAT PADG NO. 23/2024)").font = font_title
    ws3.cell(row=2, column=1, value=f"Bank: {bank.Nama_Resmi_Bank} ({bank.Masked_Alias})").font = font_subtitle

    headers_gugur = [
        "Finding ID", "Parameter", "Tanggal", "Deal ID / Seri / Tenor",
        "Nilai Realisasi", "Batas Aturan PADG", "Alasan Diskualifikasi"
    ]
    for col_idx, h_text in enumerate(headers_gugur, start=1):
        c = ws3.cell(row=4, column=col_idx, value=h_text)
        c.fill = navy_fill
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center")

    curr_row_3 = 5
    has_gugur_rows = False

    for t in tickets:
        if t.Detail_Transaksi_Gugur:
            try:
                gugur_list = json.loads(t.Detail_Transaksi_Gugur)
                if isinstance(gugur_list, list):
                    for item in gugur_list:
                        has_gugur_rows = True
                        ws3.cell(row=curr_row_3, column=1, value=t.Finding_ID).font = font_bold
                        ws3.cell(row=curr_row_3, column=2, value=t.Parameter_Code)
                        ws3.cell(row=curr_row_3, column=3, value=str(item.get("tanggal", "-")))
                        
                        detail_label = item.get("deal_id") or item.get("seri") or item.get("tenor") or "-"
                        ws3.cell(row=curr_row_3, column=4, value=str(detail_label))

                        if "spread_bps" in item:
                            val_str = f"Spread: {item['spread_bps']} bps"
                            lim_str = f"Maks: {item.get('max_allowed_bps', 20)} bps"
                        elif "tenor_hari" in item:
                            val_str = f"Tenor: {item['tenor_hari']} hari"
                            lim_str = "Min: 14 hari (2 minggu)"
                        elif "realisasi_penjualan_sekunder" in item:
                            val_str = f"Jual: Rp {item['realisasi_penjualan_sekunder']:.1f} M"
                            lim_str = f"Target: Rp {item.get('target_kewajiban_30pct', item.get('target_penjualan_sekunder_30pct', 0)):.1f} M"
                        elif "realisasi_line_du" in item:
                            val_str = f"Line DU: {item['realisasi_line_du']}, Non-DU: {item.get('realisasi_line_non_du', 0)}"
                            lim_str = "Min: 10 DU & 10 Non-DU"
                        else:
                            val_str = "-"
                            lim_str = "-"

                        ws3.cell(row=curr_row_3, column=5, value=val_str)
                        ws3.cell(row=curr_row_3, column=6, value=lim_str)
                        ws3.cell(row=curr_row_3, column=7, value=str(item.get("alasan", "-"))).font = font_fail

                        for col_idx in range(1, 8):
                            ws3.cell(row=curr_row_3, column=col_idx).border = cell_border
                        curr_row_3 += 1
            except Exception:
                pass

    if not has_gugur_rows:
        ws3.cell(row=curr_row_3, column=1, value="Tidak ada transaksi yang gugur. Seluruh transaksi memenuhi kriteria validasi PADG.").font = font_pass
        ws3.merge_cells(start_row=curr_row_3, start_column=1, end_row=curr_row_3, end_column=7)
        for col_idx in range(1, 8):
            ws3.cell(row=curr_row_3, column=col_idx).border = cell_border

    col_widths_ws3 = [24, 12, 14, 25, 20, 24, 60]
    for i, w in enumerate(col_widths_ws3, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Simpan ke stream BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
