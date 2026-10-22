"""
Evaluation Snapshot & Automated Document Generator (M.02 & M.01 Engine)
Menghasilkan dokumen resmi Bank Indonesia:
1. Memorandum M.02 (.docx): Laporan Evaluasi Triwulanan ke Pimpinan DPPK.
2. Memorandum M.01 (.docx): Pengantar Hasil Evaluasi & Rekomendasi Tindak Lanjut ke DSMM.
3. Paket 3 Lampiran Resmi (.xlsx): Multi-sheet Excel (Sebelum Threshold, Setelah Threshold, dan Tindak Lanjut).
"""

import io
import json
from datetime import datetime
from typing import Optional, Dict, Any, List
from sqlalchemy.orm import Session
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import MasterBank, EvaluationPeriod, BankScoreMatrix, FindingTicket, AuditTrail
from services.rule_engine import THRESHOLDS, WEIGHTS, PARAMETER_NAMES


# ==============================================================================
# DOCX HELPER FUNCTIONS & STYLING
# ==============================================================================
def _set_cell_background(cell, fill_hex: str):
    """Set background color of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a docx table cell in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def _set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Apply thin borders to entire table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


# ==============================================================================
# 1. GENERATOR MEMORANDUM WORD (M.02 & M.01)
# ==============================================================================
def generate_memo_docx(
    db: Session,
    memo_type: str = "m02",
    custom_narration: Optional[str] = None,
    masked: bool = False,
) -> io.BytesIO:
    """
    Generate Memorandum resmi Bank Indonesia dalam format .docx:
    - memo_type 'm02': Memorandum internal kepada Pimpinan DPPK
    - memo_type 'm01': Memorandum pengantar kepada DSMM
    """
    memo_type = memo_type.lower()
    if memo_type not in ["m02", "m01"]:
        memo_type = "m02"

    period = db.query(EvaluationPeriod).order_by(EvaluationPeriod.id.desc()).first()
    triwulan_str = period.Triwulan if period else "2026-Q2"
    stage_str = period.Active_Threshold_Stage if period else "EARLY"

    # Query 21 Bank Score Matrices
    matrices = (
        db.query(BankScoreMatrix)
        .join(MasterBank)
        .order_by(MasterBank.Bank_ID.asc())
        .all()
    )

    # Query 5 Finding Tickets
    tickets = (
        db.query(FindingTicket)
        .join(MasterBank)
        .order_by(FindingTicket.Finding_ID.asc())
        .all()
    )

    # Statistics
    total_banks = len(matrices)
    full_pass_banks = [m for m in matrices if m.compliance_status == "MEMENUHI"]
    pembinaan_banks = [m for m in matrices if m.compliance_status == "PEMBINAAN_DSMM"]
    sanksi_banks = [m for m in matrices if m.compliance_status == "SANKSI_TEGURAN"]

    doc = Document()

    # Set Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Header Bank Indonesia
    p_bi = doc.add_paragraph()
    p_bi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_bi = p_bi.add_run("BANK INDONESIA")
    run_bi.bold = True
    run_bi.font.name = "Calibri"
    run_bi.font.size = Pt(14)
    run_bi.font.color.rgb = RGBColor(15, 23, 42)

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_dept = p_dept.add_run("DEPARTEMEN PENGEMBANGAN PASAR KEUANGAN")
    run_dept.bold = True
    run_dept.font.name = "Calibri"
    run_dept.font.size = Pt(11)
    run_dept.font.color.rgb = RGBColor(71, 85, 105)

    p_memo_title = doc.add_paragraph()
    p_memo_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_memo_title = p_memo_title.add_run("MEMORANDUM")
    run_memo_title.bold = True
    run_memo_title.font.name = "Calibri"
    run_memo_title.font.size = Pt(13)
    run_memo_title.font.underline = True
    run_memo_title.font.color.rgb = RGBColor(15, 23, 42)

    # Metadata Memorandum
    table_meta = doc.add_table(rows=6, cols=3)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    memo_meta = {
        "m02": {
            "nomor": f"MEMO-{datetime.utcnow().year}/DPPK/PUVA/M.02/084",
            "kepada": "Yth. Kepala Departemen Pengembangan Pasar Keuangan (DPPK)",
            "dari": "Tim Evaluasi & Pengawasan Pasar Uang dan Pasar Valas (DPPK)",
            "perihal": f"Laporan Hasil Evaluasi Triwulanan Kewajiban Dealer Utama PUVA Periode {triwulan_str} dan Permohonan Persetujuan Hasil Rekonsiliasi",
        },
        "m01": {
            "nomor": f"MEMO-{datetime.utcnow().year}/DPPK/PUVA/M.01/085",
            "kepada": "Yth. Kepala Departemen Surveilans & Manajemen Risiko (DSMM)",
            "dari": "Kepala Departemen Pengembangan Pasar Keuangan (DPPK)",
            "perihal": f"Penyampaian Hasil Akhir Evaluasi Pemenuhan Kewajiban Dealer Utama PUVA Periode {triwulan_str} untuk Pembinaan & Pengawasan",
        },
    }

    meta_rows = [
        ("Nomor", ":", memo_meta[memo_type]["nomor"]),
        ("Kepada", ":", memo_meta[memo_type]["kepada"]),
        ("Dari", ":", memo_meta[memo_type]["dari"]),
        ("Tanggal", ":", datetime.utcnow().strftime("%d %B %Y")),
        ("Sifat", ":", "RAHASIA / SEGERA"),
        ("Perihal", ":", memo_meta[memo_type]["perihal"]),
    ]

    for idx, (lbl, sep, val) in enumerate(meta_rows):
        row = table_meta.rows[idx]
        row.cells[0].width = Inches(1.2)
        row.cells[1].width = Inches(0.2)
        row.cells[2].width = Inches(5.1)

        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.bold = True

        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(sep)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)

        p2 = row.cells[2].paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        if lbl in ["Kepada", "Nomor", "Perihal"]:
            r2.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Divider line
    p_line = doc.add_paragraph()
    p_line_run = p_line.add_run("―" * 58)
    p_line_run.font.color.rgb = RGBColor(148, 163, 184)
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ==========================================================================
    # KONTEN MEMORANDUM M.02 (KEPADA PIMPINAN DPPK)
    # ==========================================================================
    if memo_type == "m02":
        # 1. Latar Belakang
        h1 = doc.add_heading("1. Latar Belakang dan Landasan Hukum", level=2)
        h1.runs[0].font.name = "Calibri"
        h1.runs[0].font.size = Pt(11)
        h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(
            "1.1. Merujuk Peraturan Bank Indonesia (PBI) No. 6 Tahun 2024 tentang Pasar Uang dan Pasar Valuta Asing "
            "serta Peraturan Anggota Dewan Gubernur (PADG) No. 23 Tahun 2024 tentang Tata Cara Penetapan, Evaluasi, dan "
            "Kewajiban Dealer Utama Pasar Uang dan Pasar Valuta Asing (DU PUVA).\n"
            f"1.2. Telah dilaksanakan kegiatan evaluasi triwulanan atas kepatuhan 21 (dua puluh satu) Bank Dealer Utama "
            f"untuk periode {triwulan_str} (berdasarkan 60 hari kerja aktif pasar) pada tahap penetapan ambang batas: "
        )
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r_stage = p.add_run(f"{stage_str} STATE.")
        r_stage.bold = True
        r_stage.font.name = "Calibri"
        r_stage.font.size = Pt(10)

        # 2. Ringkasan Eksekutif
        h2 = doc.add_heading("2. Ringkasan Eksekutif Hasil Evaluasi & Rekonsiliasi", level=2)
        h2.runs[0].font.name = "Calibri"
        h2.runs[0].font.size = Pt(11)
        h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        # Narasi Eksekutif (Kustom dari Pengguna atau Auto-Generated)
        p_narasi = doc.add_paragraph()
        p_narasi.paragraph_format.line_spacing = 1.15
        p_narasi.paragraph_format.space_after = Pt(8)
        if custom_narration and custom_narration.strip():
            r_nar = p_narasi.add_run(custom_narration.strip())
        else:
            r_nar = p_narasi.add_run(
                f"Berdasarkan hasil kalkulasi Smart Rule Engine atas 5 sumber data primer (Bloomberg, SPPA Antasena, "
                f"Survey Counterparty Line, HARTIS, dan DS Antasena), tercatat {len(full_pass_banks)} bank memenuhi seluruh "
                f"kewajiban (100%), {len(pembinaan_banks)} bank masuk kategori pembinaan DSMM (skor 90% s.d. <100%), "
                f"dan {len(sanksi_banks)} bank dikenakan sanksi administratif teguran tertulis (skor <90%). "
                f"Seluruh proses konfirmasi dan sanggahan telah diverifikasi secara akuntabel menggunakan AI Evidence Checker "
                f"berbasis PADG No. 23 Tahun 2024."
            )
        r_nar.font.name = "Calibri"
        r_nar.font.size = Pt(10)

        # Tabel Ringkasan Sebelum vs Sesudah Rekonsiliasi
        t_sum = doc.add_table(rows=5, cols=4)
        t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(t_sum)

        headers = ["Status Evaluasi", "Sebelum Rekonsiliasi", "Pasca Rekonsiliasi", "Keterangan"]
        for c_idx, h_text in enumerate(headers):
            cell = t_sum.cell(0, c_idx)
            _set_cell_background(cell, "0F172A")
            _set_cell_margins(cell, 120, 120, 150, 150)
            p_h = cell.paragraphs[0]
            p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_h = p_h.add_run(h_text)
            r_h.bold = True
            r_h.font.color.rgb = RGBColor(255, 255, 255)
            r_h.font.name = "Calibri"
            r_h.font.size = Pt(9.5)

        sum_data = [
            ("Lolos Penuh Seluruh Kewajiban (100%)", "16 Bank", f"{len(full_pass_banks)} Bank", "Memenuhi standar PADG 23/2024"),
            ("Perlu Klarifikasi / Tiket Sanggahan", "5 Bank", "0 Bank (Selesai)", "Semua tiket diverifikasi tim BI"),
            ("Kategori Pembinaan DSMM (90 - <100%)", "4 Bank", f"{len(pembinaan_banks)} Bank", "Rekomendasi surat pembinaan"),
            ("Kategori Sanksi Teguran Tertulis (<90%)", "1 Bank", f"{len(sanksi_banks)} Bank", "Pelanggaran tenor Repo sekunder"),
        ]

        for r_idx, row_vals in enumerate(sum_data, start=1):
            for c_idx, val in enumerate(row_vals):
                cell = t_sum.cell(r_idx, c_idx)
                _set_cell_margins(cell, 80, 80, 120, 120)
                if r_idx % 2 == 1:
                    _set_cell_background(cell, "F8FAFC")
                p_c = cell.paragraphs[0]
                if c_idx in [1, 2]:
                    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_c.add_run(val)
                r_c.font.name = "Calibri"
                r_c.font.size = Pt(9)
                if c_idx == 0:
                    r_c.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 3. Rincian Tiket Sanggahan (Finding Tickets)
        h3 = doc.add_heading("3. Rincian Hasil Verifikasi Sanggahan Bank (Finding Tickets)", level=2)
        h3.runs[0].font.name = "Calibri"
        h3.runs[0].font.size = Pt(11)
        h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        t_fnd = doc.add_table(rows=len(tickets) + 1, cols=6)
        t_fnd.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(t_fnd)

        fnd_headers = ["Finding ID", "Bank DU", "Param", "Klaim Bank", "Rekomendasi AI", "Keputusan Final"]
        for c_idx, h_text in enumerate(fnd_headers):
            cell = t_fnd.cell(0, c_idx)
            _set_cell_background(cell, "0F172A")
            _set_cell_margins(cell, 120, 120, 120, 120)
            p_h = cell.paragraphs[0]
            p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_h = p_h.add_run(h_text)
            r_h.bold = True
            r_h.font.color.rgb = RGBColor(255, 255, 255)
            r_h.font.name = "Calibri"
            r_h.font.size = Pt(9)

        for idx, t in enumerate(tickets, start=1):
            bank_name = t.bank.Masked_Alias if masked else t.bank.Nama_Resmi_Bank
            vals = [
                t.Finding_ID,
                f"{t.bank.Masked_Alias}\n({bank_name})",
                t.Parameter_Code,
                t.Nilai_Klaim_Bank or "Tidak Mengajukan",
                t.AI_Recommendation_Status or "Pending",
                t.BI_Decision_Status,
            ]
            for c_idx, val in enumerate(vals):
                cell = t_fnd.cell(idx, c_idx)
                _set_cell_margins(cell, 80, 80, 100, 100)
                if idx % 2 == 1:
                    _set_cell_background(cell, "F8FAFC")
                p_c = cell.paragraphs[0]
                if c_idx in [0, 2, 4, 5]:
                    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_c.add_run(str(val))
                r_c.font.name = "Calibri"
                r_c.font.size = Pt(8.5)
                if c_idx == 5:
                    r_c.bold = True
                    if t.BI_Decision_Status == "APPROVED":
                        r_c.font.color.rgb = RGBColor(5, 150, 105)
                    elif t.BI_Decision_Status == "REJECTED":
                        r_c.font.color.rgb = RGBColor(225, 29, 72)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 4. Rekomendasi & Permohonan Persetujuan
        h4 = doc.add_heading("4. Kesimpulan dan Permohonan Persetujuan", level=2)
        h4.runs[0].font.name = "Calibri"
        h4.runs[0].font.size = Pt(11)
        h4.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        p_rek = doc.add_paragraph()
        p_rek.paragraph_format.line_spacing = 1.15
        p_rek.paragraph_format.space_after = Pt(12)
        r_rek = p_rek.add_run(
            "Sehubungan dengan selesainya proses rekonsiliasi dan verifikasi teknis atas seluruh transaksi 21 Dealer Utama PUVA, "
            "bersama ini kami memohon perkenan Bapak/Ibu Kepala Departemen Pengembangan Pasar Keuangan untuk:\n"
            f"a. Menyetujui hasil evaluasi pemenuhan kewajiban 21 Bank DU PUVA Periode {triwulan_str} sebagaimana tercantum "
            "dalam Lampiran 1, 2, dan 3;\n"
            "b. Menyetujui penyampaian Memorandum M.01 kepada Departemen Surveilans & Manajemen Risiko (DSMM) sebagai dasar "
            "pelaksanaan tindak lanjut pengawasan perbankan, pembinaan, dan penerbitan sanksi administratif sesuai PADG No. 23/2024."
        )
        r_rek.font.name = "Calibri"
        r_rek.font.size = Pt(10)

        # Lembar Tanda Tangan
        t_sign = doc.add_table(rows=3, cols=2)
        t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r in t_sign.rows:
            r.cells[0].width = Inches(3.2)
            r.cells[1].width = Inches(3.3)

        p_s0 = t_sign.cell(0, 0).paragraphs[0]
        r_s0 = p_s0.add_run("Disiapkan Oleh:\nTim Evaluasi & Rekonsiliasi PUVA\nDepartemen Pengembangan Pasar Keuangan")
        r_s0.font.name = "Calibri"
        r_s0.font.size = Pt(9.5)

        p_s1 = t_sign.cell(0, 1).paragraphs[0]
        r_s1 = p_s1.add_run("Menyetujui:\nKepala Departemen Pengembangan\nPasar Keuangan (DPPK)")
        r_s1.font.name = "Calibri"
        r_s1.font.size = Pt(9.5)
        r_s1.bold = True

        # Ruang Tanda Tangan
        t_sign.cell(1, 0).paragraphs[0].paragraph_format.space_after = Pt(40)
        t_sign.cell(1, 1).paragraphs[0].paragraph_format.space_after = Pt(40)

        p_n0 = t_sign.cell(2, 0).paragraphs[0]
        r_n0 = p_n0.add_run("Analis Senior / Manajer Tim PUVA\nNIP. 198804122012011002")
        r_n0.font.name = "Calibri"
        r_n0.font.size = Pt(9.5)
        r_n0.bold = True

        p_n1 = t_sign.cell(2, 1).paragraphs[0]
        r_n1 = p_n1.add_run("Donny Hutabarat / Pimpinan Terkait\nDirektur Eksekutif DPPK")
        r_n1.font.name = "Calibri"
        r_n1.font.size = Pt(9.5)
        r_n1.bold = True

    # ==========================================================================
    # KONTEN MEMORANDUM M.01 (PENGANTAR KE DSMM)
    # ==========================================================================
    else:
        # 1. Dasar Hukum & Maksud
        h1 = doc.add_heading("1. Dasar Hukum dan Maksud Penyampaian", level=2)
        h1.runs[0].font.name = "Calibri"
        h1.runs[0].font.size = Pt(11)
        h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(
            f"1.1. Menindaklanjuti ketentuan Pasal 28 Peraturan Anggota Dewan Gubernur No. 23 Tahun 2024 tentang "
            f"Kewajiban dan Evaluasi Dealer Utama Pasar Uang dan Pasar Valuta Asing (DU PUVA), Departemen Pengembangan "
            f"Pasar Keuangan (DPPK) menyampaikan hasil akhir evaluasi kepatuhan 21 (dua puluh satu) Bank Dealer Utama "
            f"untuk Periode {triwulan_str}.\n"
            f"1.2. Memorandum ini disampaikan kepada Departemen Surveilans & Manajemen Risiko (DSMM) sebagai rujukan resmi "
            f"dalam melaksanakan pembinaan pengawasan bank dan penegakan sanksi administratif sesuai kewenangan yang diatur dalam PADG."
        )
        r.font.name = "Calibri"
        r.font.size = Pt(10)

        # 2. Klasifikasi Hasil Evaluasi & Tindak Lanjut
        h2 = doc.add_heading("2. Klasifikasi Hasil Evaluasi & Rekomendasi Tindak Lanjut", level=2)
        h2.runs[0].font.name = "Calibri"
        h2.runs[0].font.size = Pt(11)
        h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        p_rule = doc.add_paragraph()
        p_rule.paragraph_format.line_spacing = 1.15
        p_rule.paragraph_format.space_after = Pt(8)
        r_rule = p_rule.add_run(
            "Berdasarkan perhitungan Skor Komposit Akhir (Maksimal 60% Komponen PUVA + 40% Komponen Operasi Moneter), "
            "hasil evaluasi 21 Bank DU PUVA diklasifikasikan ke dalam 3 (tiga) kategori tindak lanjut sebagai berikut:"
        )
        r_rule.font.name = "Calibri"
        r_rule.font.size = Pt(10)

        # Kategori A
        p_cat_a = doc.add_paragraph()
        p_cat_a.paragraph_format.left_indent = Inches(0.2)
        r_ca_t = p_cat_a.add_run("A. Kategori Memenuhi Seluruh Kewajiban (Skor Komposit = 100%)\n")
        r_ca_t.bold = True
        r_ca_t.font.color.rgb = RGBColor(5, 150, 105)
        names_a = [m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank for m in full_pass_banks]
        r_ca_d = p_cat_a.add_run(
            f"Sebanyak {len(full_pass_banks)} bank memenuhi seluruh kriteria tanpa catatan kelalaian: "
            f"{', '.join(names_a)}.\n"
            "Rekomendasi Tindak Lanjut: Dinyatakan LULUS evaluasi triwulanan secara penuh."
        )
        r_ca_d.font.size = Pt(9.5)

        # Kategori B
        p_cat_b = doc.add_paragraph()
        p_cat_b.paragraph_format.left_indent = Inches(0.2)
        r_cb_t = p_cat_b.add_run("B. Kategori Pembinaan Pengawasan DSMM (90.0% <= Skor Komposit < 100%)\n")
        r_cb_t.bold = True
        r_cb_t.font.color.rgb = RGBColor(217, 119, 6)
        names_b = [f"{m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank} ({m.composite_score:.1f}%)" for m in pembinaan_banks]
        r_cb_d = p_cat_b.add_run(
            f"Sebanyak {len(pembinaan_banks)} bank berada pada rentang ini: {', '.join(names_b) if names_b else 'Nihil'}.\n"
            "Rekomendasi Tindak Lanjut: Penyampaian Surat Pembinaan / Email Pembinaan oleh DSMM "
            "untuk perbaikan disiplin kuotasi pasar sekunder dan peningkatan counterparty line."
        )
        r_cb_d.font.size = Pt(9.5)

        # Kategori C
        p_cat_c = doc.add_paragraph()
        p_cat_c.paragraph_format.left_indent = Inches(0.2)
        r_cc_t = p_cat_c.add_run("C. Kategori Sanksi Administratif Teguran Tertulis (Skor Komposit < 90.0%)\n")
        r_cc_t.bold = True
        r_cc_t.font.color.rgb = RGBColor(225, 29, 72)
        names_c = [f"{m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank} ({m.composite_score:.1f}%)" for m in sanksi_banks]
        r_cc_d = p_cat_c.add_run(
            f"Sebanyak {len(sanksi_banks)} bank berada pada kategori ini: {', '.join(names_c) if names_c else 'Nihil'}.\n"
            "Rekomendasi Tindak Lanjut: Penerbitan Surat Sanksi Administratif Teguran Tertulis oleh DSMM "
            "sesuai Pasal 32 PADG No. 23 Tahun 2024 atas kelalaian transaksi pasar sekunder."
        )
        r_cc_d.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 3. Tabel Detail Bank dalam Kategori Pengawasan Khusus
        h3 = doc.add_heading("3. Ringkasan Bank Memerlukan Tindak Lanjut Khusus", level=2)
        h3.runs[0].font.name = "Calibri"
        h3.runs[0].font.size = Pt(11)
        h3.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        attention_banks = pembinaan_banks + sanksi_banks
        t_att = doc.add_table(rows=len(attention_banks) + 1, cols=5)
        t_att.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(t_att)

        att_headers = ["Bank ID", "Nama Dealer Utama", "Skor Komposit", "Parameter Perhatian", "Rekomendasi DSMM"]
        for c_idx, h_text in enumerate(att_headers):
            cell = t_att.cell(0, c_idx)
            _set_cell_background(cell, "0F172A")
            _set_cell_margins(cell, 120, 120, 120, 120)
            p_h = cell.paragraphs[0]
            p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_h = p_h.add_run(h_text)
            r_h.bold = True
            r_h.font.color.rgb = RGBColor(255, 255, 255)
            r_h.font.name = "Calibri"
            r_h.font.size = Pt(9)

        for idx, m in enumerate(attention_banks, start=1):
            bank_name = m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank
            failed_p = []
            for p_code in ["B1", "B2", "B3", "B4", "B5", "B6", "B7a", "B7b", "B8"]:
                if not getattr(m, f"{p_code.lower()}_passed", True):
                    failed_p.append(p_code)

            rec_text = "Email Pembinaan DSMM" if m.compliance_status == "PEMBINAAN_DSMM" else "Surat Teguran Tertulis"

            vals = [
                m.bank.Bank_ID,
                f"{m.bank.Masked_Alias}\n({bank_name})",
                f"{m.composite_score:.2f}%",
                ", ".join(failed_p) if failed_p else "-",
                rec_text,
            ]

            for c_idx, val in enumerate(vals):
                cell = t_att.cell(idx, c_idx)
                _set_cell_margins(cell, 80, 80, 100, 100)
                if idx % 2 == 1:
                    _set_cell_background(cell, "F8FAFC")
                p_c = cell.paragraphs[0]
                if c_idx in [0, 2, 3]:
                    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_c.add_run(str(val))
                r_c.font.name = "Calibri"
                r_c.font.size = Pt(8.5)
                if c_idx == 4:
                    r_c.bold = True
                    if "Teguran" in rec_text:
                        r_c.font.color.rgb = RGBColor(225, 29, 72)
                    else:
                        r_c.font.color.rgb = RGBColor(217, 119, 6)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # 4. Lampiran Terkait & Penutup
        p_close = doc.add_paragraph()
        p_close.paragraph_format.line_spacing = 1.15
        p_close.paragraph_format.space_after = Pt(12)
        r_cl = p_close.add_run(
            "Sebagai kelengkapan berkas tindak lanjut, bersama ini kami lampirkan:\n"
            "1. Lampiran 1: Matriks Capaian Mentah Dealer Utama (Sebelum Threshold)\n"
            "2. Lampiran 2: Matriks Skor Terbobot dan Nilai Komposit (Setelah Threshold)\n"
            "3. Lampiran 3: Keterangan Detail Hasil Evaluasi dan Rincian Temuan Per Bank\n\n"
            "Demikian kami sampaikan. Atas koordinasi dan kerja sama yang baik, kami ucapkan terima kasih."
        )
        r_cl.font.name = "Calibri"
        r_cl.font.size = Pt(10)

        # Signature Block
        t_sign = doc.add_table(rows=2, cols=1)
        t_sign.alignment = WD_TABLE_ALIGNMENT.RIGHT
        c_sign = t_sign.cell(0, 0)
        c_sign.width = Inches(3.5)
        p_s = c_sign.paragraphs[0]
        r_s = p_s.add_run("Departemen Pengembangan Pasar Keuangan\nKepala Departemen,")
        r_s.bold = True
        r_s.font.name = "Calibri"
        r_s.font.size = Pt(10)
        p_s.paragraph_format.space_after = Pt(45)

        c_name = t_sign.cell(1, 0)
        p_n = c_name.paragraphs[0]
        r_n = p_n.add_run("Donny Hutabarat\nDirektur Eksekutif")
        r_n.bold = True
        r_n.font.name = "Calibri"
        r_n.font.size = Pt(10)

    # Save to BytesIO
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_io.seek(0)
    return docx_io


# ==============================================================================
# 2. GENERATOR PAKET 3 LAMPIRAN EXCEL (LAMPIRAN 1, 2, DAN 3)
# ==============================================================================
def generate_attachments_xlsx(db: Session, masked: bool = False) -> io.BytesIO:
    """
    Menghasilkan file Excel (.xlsx) multi-sheet yang memuat 3 lampiran resmi:
    - Sheet 1: Lampiran 1 - Sebelum Threshold (Capaian Mentah B1-B8 + OM)
    - Sheet 2: Lampiran 2 - Setelah Threshold (Skor Terbobot & Komposit)
    - Sheet 3: Lampiran 3 - Tindak Lanjut (Status Kepatuhan & Rekomendasi Sanksi/Pembinaan)
    """
    period = db.query(EvaluationPeriod).order_by(EvaluationPeriod.id.desc()).first()
    triwulan_str = period.Triwulan if period else "2026-Q2"
    stage_str = period.Active_Threshold_Stage if period else "EARLY"

    matrices = (
        db.query(BankScoreMatrix)
        .join(MasterBank)
        .order_by(MasterBank.Bank_ID.asc())
        .all()
    )

    tickets = (
        db.query(FindingTicket)
        .join(MasterBank)
        .order_by(FindingTicket.Finding_ID.asc())
        .all()
    )
    tickets_by_bank: Dict[str, List[FindingTicket]] = {}
    for t in tickets:
        tickets_by_bank.setdefault(t.Bank_ID, []).append(t)

    wb = Workbook()
    wb.remove(wb.active)  # Hapus default empty sheet

    # Color Fills & Styles
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    dark_slate_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    emerald_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    light_emerald_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    light_rose_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    light_amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_title = Font(name="Calibri", size=13, bold=True, color="0F172A")
    font_subtitle = Font(name="Calibri", size=9.5, italic=True, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=9.5, color="1E293B")
    font_pass = Font(name="Calibri", size=9.5, bold=True, color="047857")
    font_fail = Font(name="Calibri", size=9.5, bold=True, color="BE123C")
    font_amber = Font(name="Calibri", size=9.5, bold=True, color="B45309")

    thin_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # --------------------------------------------------------------------------
    # SHEET 1: LAMPIRAN 1 (SEBELUM THRESHOLD - CAPAIAN MENTAH)
    # --------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="Lampiran 1 - Sebelum Thresh")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:N1")
    ws1["A1"] = "BANK INDONESIA - DEPARTEMEN PENGEMBANGAN PASAR KEUANGAN"
    ws1["A1"].font = font_subtitle

    ws1.merge_cells("A2:N2")
    ws1["A2"] = f"LAMPIRAN 1: HASIL EVALUASI REALISASI MENTAH DEALER UTAMA PUVA (SEBELUM THRESHOLD) - {triwulan_str}"
    ws1["A2"].font = font_title

    ws1.merge_cells("A3:N3")
    ws1["A3"] = f"Dasar Hukum: PBI No. 6 Tahun 2024 & PADG No. 23 Tahun 2024 | Efektif Pasar: 60 Hari Kerja | Nilai ditampilkan dalam rasio % realisasi mentah."
    ws1["A3"].font = font_subtitle

    headers_l1 = [
        "No", "Bank ID", "Nama Dealer Utama", "KBMI",
        "B1: Repo (%)", "B2: SRBI (%)", "B3: OIS (%)", "B4: Line Repo (%)",
        "B5: Line OIS (%)", "B6: Jual SRBI (%)", "B7a: Repo 14h (%)",
        "B7b: SRBI 3Bln (%)", "B8: Match OIS (%)", "Rata-rata PUVA",
    ]

    for col_idx, h_text in enumerate(headers_l1, start=1):
        cell = ws1.cell(row=5, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws1.row_dimensions[5].height = 28

    for row_idx, m in enumerate(matrices, start=6):
        bank_name = m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank
        raw_vals = [
            m.b1_raw, m.b2_raw, m.b3_raw, m.b4_raw, m.b5_raw,
            m.b6_raw, m.b7a_raw, m.b7b_raw, m.b8_raw
        ]
        avg_puva = sum(raw_vals) / len(raw_vals) if raw_vals else 0.0

        vals = [
            row_idx - 5,
            m.bank.Bank_ID,
            f"{m.bank.Masked_Alias} - {bank_name}" if not masked else m.bank.Masked_Alias,
            m.bank.Kelompok_KBMI,
            m.b1_raw / 100.0,
            m.b2_raw / 100.0,
            m.b3_raw / 100.0,
            m.b4_raw / 100.0,
            m.b5_raw / 100.0,
            m.b6_raw / 100.0,
            m.b7a_raw / 100.0,
            m.b7b_raw / 100.0,
            m.b8_raw / 100.0,
            avg_puva / 100.0,
        ]

        for col_idx, val in enumerate(vals, start=1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = cell_border
            if row_idx % 2 == 1:
                cell.fill = zebra_fill

            if col_idx in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "0.0%"

    # Auto-fit columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 11)
    ws1.column_dimensions["C"].width = 30

    # --------------------------------------------------------------------------
    # SHEET 2: LAMPIRAN 2 (SETELAH THRESHOLD - SKOR TERBOBOT & KOMPOSIT)
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Lampiran 2 - Setelah Thresh")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:Q1")
    ws2["A1"] = "BANK INDONESIA - DEPARTEMEN PENGEMBANGAN PASAR KEUANGAN"
    ws2["A1"].font = font_subtitle

    ws2.merge_cells("A2:Q2")
    ws2["A2"] = f"LAMPIRAN 2: HASIL EVALUASI SKOR TERBOBOT DEALER UTAMA PUVA (SETELAH THRESHOLD) - {triwulan_str}"
    ws2["A2"].font = font_title

    ws2.merge_cells("A3:Q3")
    ws2["A3"] = f"Tahap Threshold: {stage_str} STATE | Bobot Maksimal: PUVA 60.0% + Operasi Moneter (OM) 40.0% = Komposit 100.0%"
    ws2["A3"].font = font_subtitle

    headers_l2 = [
        "No", "Bank ID", "Nama Dealer Utama",
        "B1 (2.5%)", "B2 (2.5%)", "B3 (2.5%)", "B4 (7.5%)", "B5 (2.5%)",
        "B6 (7.5%)", "B7a (20%)", "B7b (5%)", "B8 (10%)", "PUVA (60%)",
        "OM1 (15%)", "OM2 (20%)", "OM3 (5%)", "Komposit (100%)",
    ]

    for col_idx, h_text in enumerate(headers_l2, start=1):
        cell = ws2.cell(row=5, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = dark_slate_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws2.row_dimensions[5].height = 28

    for row_idx, m in enumerate(matrices, start=6):
        bank_name = m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank
        vals = [
            row_idx - 5,
            m.bank.Bank_ID,
            f"{m.bank.Masked_Alias} - {bank_name}" if not masked else m.bank.Masked_Alias,
            m.b1_score / 100.0,
            m.b2_score / 100.0,
            m.b3_score / 100.0,
            m.b4_score / 100.0,
            m.b5_score / 100.0,
            m.b6_score / 100.0,
            m.b7a_score / 100.0,
            m.b7b_score / 100.0,
            m.b8_score / 100.0,
            m.total_puva_score / 100.0,
            m.om1_score / 100.0,
            m.om2_score / 100.0,
            m.om3_score / 100.0,
            m.composite_score / 100.0,
        ]

        for col_idx, val in enumerate(vals, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = cell_border
            if row_idx % 2 == 1:
                cell.fill = zebra_fill

            if col_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in [13, 17]:
                # Kolom Subtotal PUVA dan Komposit Total
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = font_bold
                cell.number_format = "0.0%"
                if col_idx == 17:
                    if val >= 1.0:
                        cell.fill = light_emerald_fill
                        cell.font = font_pass
                    elif val >= 0.90:
                        cell.fill = light_amber_fill
                        cell.font = font_amber
                    else:
                        cell.fill = light_rose_fill
                        cell.font = font_fail
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "0.0%"

    # Auto-fit columns
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 11)
    ws2.column_dimensions["C"].width = 28

    # --------------------------------------------------------------------------
    # SHEET 3: LAMPIRAN 3 (TINDAK LANJUT & STATUS KEPATUHAN PER BANK)
    # --------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Lampiran 3 - Tindak Lanjut")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:I1")
    ws3["A1"] = "BANK INDONESIA - DEPARTEMEN PENGEMBANGAN PASAR KEUANGAN"
    ws3["A1"].font = font_subtitle

    ws3.merge_cells("A2:I2")
    ws3["A2"] = f"LAMPIRAN 3: KETERANGAN HASIL EVALUASI & REKOMENDASI TINDAK LANJUT PER BANK - {triwulan_str}"
    ws3["A2"].font = font_title

    ws3.merge_cells("A3:I3")
    ws3["A3"] = "Rujukan Tindak Lanjut: Skor 100% (Lulus Penuh) | 90% s.d. <100% (Email Pembinaan DSMM) | <90% (Sanksi Administratif Teguran Tertulis)"
    ws3["A3"].font = font_subtitle

    headers_l3 = [
        "No", "Bank ID", "Nama Dealer Utama", "KBMI",
        "Skor Komposit", "Status Kepatuhan", "Parameter Perhatian",
        "Tiket Finding Sanggahan", "Rekomendasi Tindak Lanjut",
    ]

    for col_idx, h_text in enumerate(headers_l3, start=1):
        cell = ws3.cell(row=5, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws3.row_dimensions[5].height = 28

    for row_idx, m in enumerate(matrices, start=6):
        bank_name = m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank

        # Parameter yang gugur
        failed_params = []
        for p_code in ["B1", "B2", "B3", "B4", "B5", "B6", "B7a", "B7b", "B8"]:
            if not getattr(m, f"{p_code.lower()}_passed", True):
                failed_params.append(p_code)

        bank_tickets = tickets_by_bank.get(m.Bank_ID, [])
        ticket_str = ", ".join([f"{t.Finding_ID} ({t.BI_Decision_Status})" for t in bank_tickets]) if bank_tickets else "-"

        if m.compliance_status == "MEMENUHI":
            tindak_lanjut = "Lulus Penuh (Memenuhi Seluruh Kewajiban DU PUVA)"
        elif m.compliance_status == "PEMBINAAN_DSMM":
            tindak_lanjut = "Penyampaian Surat / Email Pembinaan oleh DSMM"
        else:
            tindak_lanjut = "Penerbitan Surat Sanksi Administratif Teguran Tertulis oleh DSMM"

        vals = [
            row_idx - 5,
            m.bank.Bank_ID,
            f"{m.bank.Masked_Alias} - {bank_name}" if not masked else m.bank.Masked_Alias,
            m.bank.Kelompok_KBMI,
            m.composite_score / 100.0,
            m.compliance_status,
            ", ".join(failed_params) if failed_params else "Nihil (Lolos Semua)",
            ticket_str,
            tindak_lanjut,
        ]

        for col_idx, val in enumerate(vals, start=1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = cell_border
            if row_idx % 2 == 1:
                cell.fill = zebra_fill

            if col_idx in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [3, 7, 8]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "0.0%"
                cell.font = font_bold
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_bold
                if val == "MEMENUHI":
                    cell.font = font_pass
                    cell.fill = light_emerald_fill
                elif val == "PEMBINAAN_DSMM":
                    cell.font = font_amber
                    cell.fill = light_amber_fill
                else:
                    cell.font = font_fail
                    cell.fill = light_rose_fill
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = font_bold
                if "Teguran" in val:
                    cell.font = font_fail
                elif "Pembinaan" in val:
                    cell.font = font_amber
                else:
                    cell.font = font_pass

    # Auto-fit columns
    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 11)
    ws3.column_dimensions["C"].width = 28
    ws3.column_dimensions["I"].width = 42

    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)
    return excel_io


# ==============================================================================
# 3. GENERATOR STRUCTURED PREVIEW JSON (UNTUK LIVE PREVIEW FRONTEND)
# ==============================================================================
def get_memo_preview_data(db: Session, memo_type: str = "m02", masked: bool = False) -> Dict[str, Any]:
    """
    Menghasilkan data struktural memorandum untuk ditampilkan secara real-time
    pada Live Preview di UI sebelum diunduh dalam format .docx atau .xlsx.
    """
    memo_type = memo_type.lower()
    period = db.query(EvaluationPeriod).order_by(EvaluationPeriod.id.desc()).first()
    triwulan_str = period.Triwulan if period else "2026-Q2"
    stage_str = period.Active_Threshold_Stage if period else "EARLY"

    matrices = (
        db.query(BankScoreMatrix)
        .join(MasterBank)
        .order_by(MasterBank.Bank_ID.asc())
        .all()
    )

    tickets = (
        db.query(FindingTicket)
        .join(MasterBank)
        .order_by(FindingTicket.Finding_ID.asc())
        .all()
    )

    full_pass_banks = [m for m in matrices if m.compliance_status == "MEMENUHI"]
    pembinaan_banks = [m for m in matrices if m.compliance_status == "PEMBINAAN_DSMM"]
    sanksi_banks = [m for m in matrices if m.compliance_status == "SANKSI_TEGURAN"]

    if memo_type == "m02":
        return {
            "type": "M.02",
            "title": "MEMORANDUM LAPORAN EVALUASI TRIWULANAN (M.02)",
            "nomor": f"MEMO-{datetime.utcnow().year}/DPPK/PUVA/M.02/084",
            "kepada": "Yth. Kepala Departemen Pengembangan Pasar Keuangan (DPPK)",
            "dari": "Tim Evaluasi & Pengawasan Pasar Uang dan Pasar Valas (DPPK)",
            "tanggal": datetime.utcnow().strftime("%d %B %Y"),
            "sifat": "RAHASIA / SEGERA",
            "lampiran": "3 (tiga) berkas matriks evaluasi",
            "perihal": f"Laporan Hasil Evaluasi Triwulanan Kewajiban Dealer Utama PUVA Periode {triwulan_str} dan Permohonan Persetujuan Hasil Rekonsiliasi",
            "period": triwulan_str,
            "stage": stage_str,
            "default_narration": (
                f"Berdasarkan hasil kalkulasi Smart Rule Engine atas 5 sumber data primer (Bloomberg, SPPA Antasena, "
                f"Survey Counterparty Line, HARTIS, dan DS Antasena) untuk periode {triwulan_str} (60 hari kerja), "
                f"tercatat {len(full_pass_banks)} bank memenuhi seluruh kewajiban (100%), {len(pembinaan_banks)} bank masuk "
                f"kategori pembinaan DSMM (skor 90% s.d. <100%), dan {len(sanksi_banks)} bank dikenakan sanksi administratif "
                f"teguran tertulis (skor <90%). Seluruh tiket sanggahan telah diverifikasi melalui AI Evidence Checker berbasis PADG No. 23/2024."
            ),
            "summary_stats": {
                "total_banks": len(matrices),
                "full_pass": len(full_pass_banks),
                "pembinaan": len(pembinaan_banks),
                "sanksi": len(sanksi_banks),
            },
            "findings": [
                {
                    "finding_id": t.Finding_ID,
                    "bank_id": t.Bank_ID,
                    "bank_name": t.bank.Masked_Alias if masked else t.bank.Nama_Resmi_Bank,
                    "parameter": t.Parameter_Code,
                    "nilai_bi": t.Nilai_Rekap_BI,
                    "klaim_bank": t.Nilai_Klaim_Bank or "Belum diajukan",
                    "ai_recom": t.AI_Recommendation_Status or "PENDING",
                    "decision": t.BI_Decision_Status,
                }
                for t in tickets
            ],
        }
    else:
        return {
            "type": "M.01",
            "title": "MEMORANDUM PENGANTAR HASIL EVALUASI KE DSMM (M.01)",
            "nomor": f"MEMO-{datetime.utcnow().year}/DPPK/PUVA/M.01/085",
            "kepada": "Yth. Kepala Departemen Surveilans & Manajemen Risiko (DSMM)",
            "dari": "Kepala Departemen Pengembangan Pasar Keuangan (DPPK)",
            "tanggal": datetime.utcnow().strftime("%d %B %Y"),
            "sifat": "RAHASIA / SEGERA",
            "lampiran": "3 (tiga) berkas matriks evaluasi",
            "perihal": f"Penyampaian Hasil Akhir Evaluasi Pemenuhan Kewajiban Dealer Utama PUVA Periode {triwulan_str} untuk Pembinaan & Pengawasan Perbankan",
            "period": triwulan_str,
            "stage": stage_str,
            "default_narration": (
                f"Menindaklanjuti ketentuan Pasal 28 PADG No. 23 Tahun 2024, Departemen Pengembangan Pasar Keuangan (DPPK) "
                f"menyampaikan hasil akhir evaluasi 21 Dealer Utama PUVA periode {triwulan_str} pasca-rekonsiliasi. "
                f"Hasil evaluasi memuat daftar bank yang memenuhi kewajiban, bank yang memerlukan email pembinaan DSMM, "
                f"serta bank yang dikenakan sanksi administratif teguran tertulis."
            ),
            "categories": {
                "lulus": [m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank for m in full_pass_banks],
                "pembinaan": [
                    {
                        "bank_name": m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank,
                        "score": m.composite_score,
                        "notes": "Di bawah ambang batas B2 (SRBI) / B3 (OIS) / B4 (Line Repo)",
                    }
                    for m in pembinaan_banks
                ],
                "sanksi": [
                    {
                        "bank_name": m.bank.Masked_Alias if masked else m.bank.Nama_Resmi_Bank,
                        "score": m.composite_score,
                        "notes": "Pelanggaran tenor transaksi Repo sekunder (< 14 hari) pada B7a",
                    }
                    for m in sanksi_banks
                ],
            },
        }
